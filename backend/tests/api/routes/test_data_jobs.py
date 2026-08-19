import io
import zipfile
from datetime import date, timedelta
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlmodel import Session, select

from app.core.config import settings
from app.data_jobs import cleanup_expired_data_jobs, recover_interrupted_data_jobs
from app.models import (
    Budget,
    Category,
    DataJob,
    DataJobError,
    DataJobStatus,
    DataJobType,
    Transaction,
    get_datetime_utc,
)
from tests.utils.user import create_random_user


def _build_import_workbook(owner_login_name: str, handler_login_name: str) -> bytes:
    workbook = Workbook()
    readme = workbook.active
    readme.title = "README"
    readme.append(["Template Version", "v2"])
    readme.append(["Sheet", "Description"])

    categories_sheet = workbook.create_sheet("Categories")
    categories_sheet.append(["owner_login_name", "name", "parent_name", "color"])
    categories_sheet.append([owner_login_name, "Groceries", "", "#22c55e"])
    categories_sheet.append([owner_login_name, "Fruit", "Groceries", "#16a34a"])

    budgets_sheet = workbook.create_sheet("Budgets")
    budgets_sheet.append(["owner_login_name", "name", "year", "period", "amount"])
    budgets_sheet.append([owner_login_name, "Food Budget", 2026, "MONTH", 1000])

    transactions_sheet = workbook.create_sheet("Transactions")
    transactions_sheet.append(
        [
            "owner_login_name",
            "transaction_date",
            "transaction_type",
            "category_name",
            "budget_name",
            "amount",
            "entry_status",
            "handler_login_name",
            "summary",
            "detail_note",
            "detail_items",
        ]
    )
    transactions_sheet.append(
        [
            owner_login_name,
            "2026-05-30",
            "EXPENSE",
            "Groceries",
            "Food Budget",
            32.5,
            "ENTERED",
            handler_login_name,
            "逛超市",
            "周末采购",
            "牛奶\n鸡蛋",
        ]
    )

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_download_data_job_template(
    client: TestClient, superuser_token_headers: dict[str, str]
) -> None:
    response = client.get(
        f"{settings.API_V1_STR}/system/data-jobs/template",
        headers=superuser_token_headers,
    )
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content), read_only=True)
    assert workbook.sheetnames == ["README", "Categories", "Budgets", "Transactions"]
    assert workbook["README"]["B1"].value == "v2"


def test_export_job_creates_downloadable_excel(
    client: TestClient, superuser_token_headers: dict[str, str], db: Session
) -> None:
    owner = create_random_user(db)
    category = Category(owner_id=owner.id, name="Travel", color="#123456")
    db.add(category)
    db.commit()
    db.refresh(category)
    budget = Budget(
        owner_id=owner.id, name="Trip", year=2026, period=1, amount_cents=50000
    )
    db.add(budget)
    db.commit()
    db.refresh(budget)
    transaction = Transaction(
        owner_id=owner.id,
        category_id=category.id,
        budget_id=budget.id,
        transaction_date=date(2026, 5, 30),
        transaction_type=2,
        amount_cents=4200,
        entry_status=3,
        handler_user_id=owner.id,
        handler_name=owner.login_name,
        summary="打车",
        description="打车",
        detail={"note": "机场回家", "items": [{"name": "出租车"}]},
    )
    db.add(transaction)
    db.commit()

    create_response = client.post(
        f"{settings.API_V1_STR}/system/data-jobs/export",
        headers=superuser_token_headers,
    )
    assert create_response.status_code == 200
    payload = create_response.json()
    assert payload["status"] == "PENDING"

    db.expire_all()
    job = db.exec(select(DataJob).order_by(DataJob.created_at.desc())).first()
    assert job is not None
    db.refresh(job)
    assert job.status == "SUCCEEDED"
    assert job.result_file_path

    download_response = client.get(
        f"{settings.API_V1_STR}/system/data-jobs/{job.id}/result",
        headers=superuser_token_headers,
    )
    assert download_response.status_code == 200
    workbook = load_workbook(io.BytesIO(download_response.content), read_only=True)
    assert "Transactions" in workbook.sheetnames
    transaction_rows = list(workbook["Transactions"].iter_rows(values_only=True))
    headers = transaction_rows[0]
    summary_index = headers.index("summary")
    detail_note_index = headers.index("detail_note")
    detail_items_index = headers.index("detail_items")
    assert any(
        row[summary_index] == "打车"
        and row[detail_note_index] == "机场回家"
        and "出租车" in (row[detail_items_index] or "")
        for row in transaction_rows[1:]
    )


def test_import_job_creates_records_and_error_file(
    client: TestClient, superuser_token_headers: dict[str, str], db: Session
) -> None:
    owner = create_random_user(db)
    handler = create_random_user(db)
    workbook_bytes = _build_import_workbook(owner.login_name, handler.login_name)

    response = client.post(
        f"{settings.API_V1_STR}/system/data-jobs/import",
        headers=superuser_token_headers,
        files={
            "file": (
                "homefin-import.xlsx",
                workbook_bytes,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 200

    db.expire_all()
    job = db.exec(select(DataJob).order_by(DataJob.created_at.desc())).first()
    assert job is not None
    db.refresh(job)
    assert job.status == "SUCCEEDED"
    assert job.success_rows == 4

    imported_category = db.exec(
        select(Category).where(
            Category.owner_id == owner.id, Category.name == "Groceries"
        )
    ).first()
    assert imported_category is not None

    imported_budget = db.exec(
        select(Budget).where(Budget.owner_id == owner.id, Budget.name == "Food Budget")
    ).first()
    assert imported_budget is not None

    imported_transaction = db.exec(
        select(Transaction).where(
            Transaction.owner_id == owner.id, Transaction.summary == "逛超市"
        )
    ).first()
    assert imported_transaction is not None
    assert imported_transaction.handler_user_id == handler.id
    assert imported_transaction.description == "逛超市"
    assert imported_transaction.detail == {
        "note": "周末采购",
        "items": [{"name": "牛奶"}, {"name": "鸡蛋"}],
    }

    errors = db.exec(select(DataJobError).where(DataJobError.job_id == job.id)).all()
    assert errors == []


def test_import_job_generates_error_file_for_invalid_rows(
    client: TestClient, superuser_token_headers: dict[str, str], db: Session
) -> None:
    owner = create_random_user(db)
    workbook = Workbook()
    readme = workbook.active
    readme.title = "README"
    readme.append(["Template Version", "v2"])
    readme.append(["Sheet", "Description"])

    categories_sheet = workbook.create_sheet("Categories")
    categories_sheet.append(["owner_login_name", "name", "parent_name", "color"])

    budgets_sheet = workbook.create_sheet("Budgets")
    budgets_sheet.append(["owner_login_name", "name", "year", "period", "amount"])

    transactions_sheet = workbook.create_sheet("Transactions")
    transactions_sheet.append(
        [
            "owner_login_name",
            "transaction_date",
            "transaction_type",
            "category_name",
            "budget_name",
            "amount",
            "entry_status",
            "handler_login_name",
            "summary",
            "detail_note",
            "detail_items",
        ]
    )
    transactions_sheet.append(
        [
            owner.login_name,
            "2026-05-30",
            "EXPENSE",
            "Missing",
            "",
            12,
            "ENTERED",
            "",
            "坏数据",
            "",
            "",
        ]
    )
    buffer = io.BytesIO()
    workbook.save(buffer)

    response = client.post(
        f"{settings.API_V1_STR}/system/data-jobs/import",
        headers=superuser_token_headers,
        files={
            "file": (
                "homefin-import.xlsx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 200

    job = db.exec(select(DataJob).order_by(DataJob.created_at.desc())).first()
    assert job is not None
    db.refresh(job)
    assert job.status == "PARTIAL_SUCCESS"
    assert job.error_file_path

    download_response = client.get(
        f"{settings.API_V1_STR}/system/data-jobs/{job.id}/errors",
        headers=superuser_token_headers,
    )
    assert download_response.status_code == 200
    workbook = load_workbook(io.BytesIO(download_response.content), read_only=True)
    rows = list(workbook["Errors"].iter_rows(values_only=True))
    assert any(row[3] == "Category not found" for row in rows[1:])


def test_import_rejects_oversized_and_invalid_archives(
    client: TestClient, superuser_token_headers: dict[str, str]
) -> None:
    with patch("app.core.config.settings.DATA_JOB_MAX_UPLOAD_BYTES", 4):
        oversized = client.post(
            f"{settings.API_V1_STR}/system/data-jobs/import",
            headers=superuser_token_headers,
            files={"file": ("large.xlsx", b"12345", "application/octet-stream")},
        )
    assert oversized.status_code == 413

    invalid = client.post(
        f"{settings.API_V1_STR}/system/data-jobs/import",
        headers=superuser_token_headers,
        files={"file": ("invalid.xlsx", b"not-a-zip", "application/octet-stream")},
    )
    assert invalid.status_code == 400
    assert invalid.json() == {"detail": "Uploaded file is not a valid .xlsx archive"}


def test_import_rejects_archive_safety_limit_and_removes_upload(
    client: TestClient,
    superuser_token_headers: dict[str, str],
    tmp_path: Path,
) -> None:
    archive_bytes = io.BytesIO()
    with zipfile.ZipFile(
        archive_bytes, mode="w", compression=zipfile.ZIP_DEFLATED
    ) as archive:
        archive.writestr("first.xml", "first")
        archive.writestr("second.xml", "second")

    with (
        patch("app.data_jobs.JOB_STORAGE_DIR", tmp_path),
        patch("app.core.config.settings.DATA_JOB_MAX_ARCHIVE_FILES", 1),
    ):
        response = client.post(
            f"{settings.API_V1_STR}/system/data-jobs/import",
            headers=superuser_token_headers,
            files={
                "file": (
                    "too-many-members.xlsx",
                    archive_bytes.getvalue(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

    assert response.status_code == 400
    assert response.json() == {"detail": "Excel archive contains too many files"}
    assert list(tmp_path.iterdir()) == []


def test_export_rejects_when_active_job_limit_is_reached(
    client: TestClient,
    superuser_token_headers: dict[str, str],
    db: Session,
) -> None:
    owner = create_random_user(db)
    active_job = DataJob(
        job_type=DataJobType.EXPORT,
        status=DataJobStatus.PROCESSING,
        created_by=owner.id,
    )
    db.add(active_job)
    db.commit()

    with patch("app.core.config.settings.DATA_JOB_MAX_ACTIVE_JOBS", 1):
        response = client.post(
            f"{settings.API_V1_STR}/system/data-jobs/export",
            headers=superuser_token_headers,
        )

    assert response.status_code == 429
    assert response.json() == {"detail": "Too many data jobs are already running"}
    db.delete(active_job)
    db.commit()


def test_data_job_recovery_and_retention_cleanup(
    db: Session,
    tmp_path: Path,
) -> None:
    owner = create_random_user(db)
    now = get_datetime_utc()
    stale_job = DataJob(
        job_type=DataJobType.EXPORT,
        status=DataJobStatus.PROCESSING,
        created_by=owner.id,
        created_at=now - timedelta(hours=2),
        started_at=now - timedelta(hours=2),
    )
    expired_file = tmp_path / "expired.xlsx"
    expired_file.write_bytes(b"expired")
    outside_file = tmp_path.parent / f"{tmp_path.name}-outside.xlsx"
    outside_file.write_bytes(b"outside")
    expired_job = DataJob(
        job_type=DataJobType.IMPORT,
        status=DataJobStatus.FAILED,
        created_by=owner.id,
        created_at=now - timedelta(days=31),
        source_file_path=str(expired_file),
        error_file_path=str(outside_file),
    )
    db.add(stale_job)
    db.add(expired_job)
    db.commit()
    stale_job_id = stale_job.id
    expired_job_id = expired_job.id

    with (
        patch("app.data_jobs.JOB_STORAGE_DIR", tmp_path),
        patch("app.core.config.settings.DATA_JOB_STALE_MINUTES", 60),
        patch("app.core.config.settings.DATA_JOB_RETENTION_DAYS", 30),
    ):
        assert recover_interrupted_data_jobs() == 1
        assert cleanup_expired_data_jobs() == 1

    db.expire_all()
    recovered = db.get(DataJob, stale_job_id)
    assert recovered is not None
    assert recovered.status == DataJobStatus.FAILED
    assert recovered.failure_reason == "Job interrupted by application restart"
    assert recovered.finished_at is not None
    assert db.get(DataJob, expired_job_id) is None
    assert not expired_file.exists()
    assert outside_file.read_bytes() == b"outside"
    outside_file.unlink()
