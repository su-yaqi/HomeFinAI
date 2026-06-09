from __future__ import annotations

import io
import uuid
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from sqlmodel import Session, col, select

from app.core.config import settings
from app.core.db import engine
from app.models import (
    Budget,
    BudgetPeriod,
    Category,
    DataJob,
    DataJobError,
    DataJobPublic,
    DataJobStatus,
    DataJobType,
    EntryStatus,
    Transaction,
    TransactionType,
    User,
)
from app.utils import amount_to_cents, cents_to_amount, logger

TEMPLATE_VERSION = "v2"
README_SHEET = "README"
CATEGORIES_SHEET = "Categories"
BUDGETS_SHEET = "Budgets"
TRANSACTIONS_SHEET = "Transactions"
JOB_BATCH_SIZE = 200
JOB_STORAGE_DIR = Path(settings.DATA_JOBS_STORAGE_DIR)

TRANSACTION_HEADERS = [
    "owner_login_name",
    "transaction_date",
    "transaction_type",
    "category_name",
    "budget_name",
    "amount",
    "entry_status",
    "handler_login_name",
    "summary",
    "detail_note",
    "detail_items",
]


class ImportValidationError(Exception):
    def __init__(
        self,
        *,
        sheet_name: str,
        row_number: int,
        field_name: str,
        message: str,
        raw_key: str | None = None,
    ) -> None:
        super().__init__(message)
        self.sheet_name = sheet_name
        self.row_number = row_number
        self.field_name = field_name
        self.message = message
        self.raw_key = raw_key


@dataclass
class PendingTransactionRow:
    transaction: Transaction
    row_number: int
    raw_key: str | None


def ensure_job_storage_dir() -> Path:
    JOB_STORAGE_DIR.mkdir(parents=True, exist_ok=True)
    return JOB_STORAGE_DIR


def data_job_public(job: DataJob) -> DataJobPublic:
    return DataJobPublic.model_validate(
        {
            **job.model_dump(),
            "has_result_file": bool(job.result_file_path),
            "has_error_file": bool(job.error_file_path),
        }
    )


def create_template_workbook_bytes() -> bytes:
    workbook = Workbook()
    readme = workbook.active
    readme.title = README_SHEET
    _populate_readme_sheet(readme)

    categories_sheet = workbook.create_sheet(CATEGORIES_SHEET)
    categories_sheet.append(["owner_login_name", "name", "parent_name", "color"])
    categories_sheet.append(["alice", "Groceries", "", "#22c55e"])
    categories_sheet.append(["alice", "Fruit", "Groceries", "#16a34a"])

    budgets_sheet = workbook.create_sheet(BUDGETS_SHEET)
    budgets_sheet.append(["owner_login_name", "name", "year", "period", "amount"])
    budgets_sheet.append(["alice", "Food Budget", 2026, "MONTH", 1500.00])

    transactions_sheet = workbook.create_sheet(TRANSACTIONS_SHEET)
    transactions_sheet.append(TRANSACTION_HEADERS)
    transactions_sheet.append(
        [
            "alice",
            "2026-05-30",
            "EXPENSE",
            "Groceries",
            "Food Budget",
            86.50,
            "ENTERED",
            "alice",
            "逛超市",
            "周末采购",
            "牛奶\n鸡蛋",
        ]
    )

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def save_template_workbook(path: Path) -> None:
    path.write_bytes(create_template_workbook_bytes())


def save_uploaded_file(upload_bytes: bytes, suffix: str = ".xlsx") -> Path:
    storage_dir = ensure_job_storage_dir()
    file_path = storage_dir / f"{uuid.uuid4()}{suffix}"
    file_path.write_bytes(upload_bytes)
    return file_path


def process_job(job_id: uuid.UUID) -> None:
    with Session(engine) as session:
        job = session.get(DataJob, job_id)
        if not job:
            return
        job_type = job.job_type
        job.status = DataJobStatus.PROCESSING
        job.progress_percent = 0
        job.started_at = job.started_at or _utcnow()
        job.failure_reason = None
        session.add(job)
        session.commit()

    try:
        if job_type == DataJobType.EXPORT:
            _process_export_job(job_id)
        else:
            _process_import_job(job_id)
    except Exception as exc:  # noqa: BLE001
        logger.exception("Data job %s failed", job_id)
        _mark_job_failed(job_id, str(exc))


def _process_export_job(job_id: uuid.UUID) -> None:
    with Session(engine) as session:
        job = session.get(DataJob, job_id)
        if not job:
            return

        storage_dir = ensure_job_storage_dir()
        result_path = storage_dir / f"export-{job.id}.xlsx"
        workbook = Workbook(write_only=True)

        readme = workbook.create_sheet(README_SHEET)
        _populate_readme_sheet(readme)

        users = session.exec(select(User)).all()
        users_by_id = {user.id: user for user in users}

        categories = session.exec(select(Category).order_by(Category.owner_id, Category.name)).all()
        category_by_id = {category.id: category for category in categories}
        categories_sheet = workbook.create_sheet(CATEGORIES_SHEET)
        categories_sheet.append(["owner_login_name", "name", "parent_name", "color"])
        for category in categories:
            owner = users_by_id.get(category.owner_id)
            parent_name = (
                category_by_id[category.parent_id].name
                if category.parent_id and category.parent_id in category_by_id
                else ""
            )
            categories_sheet.append(
                [
                    owner.login_name if owner else "",
                    category.name,
                    parent_name,
                    category.color,
                ]
            )

        budgets = session.exec(
            select(Budget).order_by(
                col(Budget.owner_id), col(Budget.year).desc(), col(Budget.created_at).desc()
            )
        ).all()
        budgets_by_id = {budget.id: budget for budget in budgets}
        budgets_sheet = workbook.create_sheet(BUDGETS_SHEET)
        budgets_sheet.append(["owner_login_name", "name", "year", "period", "amount"])
        for budget in budgets:
            owner = users_by_id.get(budget.owner_id)
            budgets_sheet.append(
                [
                    owner.login_name if owner else "",
                    budget.name,
                    budget.year,
                    BudgetPeriod(budget.period).name,
                    cents_to_amount(budget.amount_cents),
                ]
            )

        transactions_sheet = workbook.create_sheet(TRANSACTIONS_SHEET)
        transactions_sheet.append(TRANSACTION_HEADERS)

        total_rows = len(categories) + len(budgets)
        offset = 0
        processed = 0
        while True:
            transactions = session.exec(
                select(Transaction)
                .order_by(
                    col(Transaction.transaction_date).desc(),
                    col(Transaction.created_at).desc(),
                )
                .offset(offset)
                .limit(JOB_BATCH_SIZE)
            ).all()
            if not transactions:
                break
            total_rows += len(transactions)
            for transaction in transactions:
                owner = users_by_id.get(transaction.owner_id)
                category = category_by_id.get(transaction.category_id)
                budget = (
                    budgets_by_id.get(transaction.budget_id)
                    if transaction.budget_id is not None
                    else None
                )
                handler_user = (
                    users_by_id.get(transaction.handler_user_id)
                    if transaction.handler_user_id is not None
                    else None
                )
                transactions_sheet.append(
                    [
                        owner.login_name if owner else "",
                        transaction.transaction_date.isoformat(),
                        TransactionType(transaction.transaction_type).name,
                        category.name if category else "",
                        budget.name if budget else "",
                        cents_to_amount(transaction.amount_cents),
                        EntryStatus(transaction.entry_status).name,
                        handler_user.login_name if handler_user else "",
                        transaction.summary or transaction.description or "",
                        _detail_note(transaction.detail),
                        _serialize_detail_items(transaction.detail),
                    ]
                )
                processed += 1
            offset += len(transactions)
            _update_job_progress(
                session,
                job,
                progress_percent=min(95, 30 + processed // max(1, JOB_BATCH_SIZE)),
                total_rows=total_rows,
            )

        workbook.save(result_path)
        job.result_file_path = str(result_path)
        job.total_rows = total_rows
        job.success_rows = total_rows
        job.progress_percent = 100
        job.status = DataJobStatus.SUCCEEDED
        job.finished_at = _utcnow()
        session.add(job)
        session.commit()


def _process_import_job(job_id: uuid.UUID) -> None:
    with Session(engine) as session:
        job = session.get(DataJob, job_id)
        if not job or not job.source_file_path:
            raise RuntimeError("Import file not found")

        workbook = load_workbook(job.source_file_path, read_only=True, data_only=True)
        _validate_required_sheets(workbook.sheetnames)
        template_version = _read_template_version(workbook[README_SHEET])
        if template_version != TEMPLATE_VERSION:
            raise RuntimeError(
                f"Unsupported template version: {template_version or 'unknown'}"
            )
        job.template_version = template_version

        users = session.exec(select(User)).all()
        users_by_login = {user.login_name: user for user in users}
        categories_by_key = {
            (category.owner_id, category.name): category
            for category in session.exec(select(Category)).all()
        }
        budgets_by_key = {
            (budget.owner_id, budget.name, budget.year, budget.period): budget
            for budget in session.exec(select(Budget)).all()
        }

        total_rows = sum(
            max(workbook[sheet_name].max_row - 1, 0)
            for sheet_name in [CATEGORIES_SHEET, BUDGETS_SHEET, TRANSACTIONS_SHEET]
        )
        job.total_rows = total_rows
        session.add(job)
        session.commit()

        processed = 0

        for row_number, row in _iter_sheet_rows(
            workbook[CATEGORIES_SHEET],
            ["owner_login_name", "name", "parent_name", "color"],
        ):
            try:
                owner = _require_user(users_by_login, row["owner_login_name"], CATEGORIES_SHEET, row_number)
                name = _require_text(row, "name", CATEGORIES_SHEET, row_number)
                color = _require_text(row, "color", CATEGORIES_SHEET, row_number)
                parent_name = _optional_text(row.get("parent_name"))
                if not color.startswith("#") or len(color) not in {4, 7}:
                    raise ImportValidationError(
                        sheet_name=CATEGORIES_SHEET,
                        row_number=row_number,
                        field_name="color",
                        message="Invalid color",
                        raw_key=name,
                    )
                category_key = (owner.id, name)
                if category_key in categories_by_key:
                    raise ImportValidationError(
                        sheet_name=CATEGORIES_SHEET,
                        row_number=row_number,
                        field_name="name",
                        message="Category already exists",
                        raw_key=name,
                    )
                parent = categories_by_key.get((owner.id, parent_name)) if parent_name else None
                if parent_name and parent is None:
                    raise ImportValidationError(
                        sheet_name=CATEGORIES_SHEET,
                        row_number=row_number,
                        field_name="parent_name",
                        message="Parent category not found",
                        raw_key=name,
                    )
                category = Category(
                    owner_id=owner.id,
                    name=name,
                    parent_id=parent.id if parent else None,
                    color=color,
                )
                session.add(category)
                session.commit()
                session.refresh(category)
                categories_by_key[category_key] = category
                job.success_rows += 1
            except ImportValidationError as exc:
                _record_job_error(session, job.id, exc)
                job.failed_rows += 1
            processed += 1
            _update_job_progress(
                session,
                job,
                progress_percent=_progress_percent(processed, total_rows),
            )

        for row_number, row in _iter_sheet_rows(
            workbook[BUDGETS_SHEET],
            ["owner_login_name", "name", "year", "period", "amount"],
        ):
            try:
                owner = _require_user(users_by_login, row["owner_login_name"], BUDGETS_SHEET, row_number)
                name = _require_text(row, "name", BUDGETS_SHEET, row_number)
                year = _require_int(row, "year", BUDGETS_SHEET, row_number)
                period = _parse_budget_period(row.get("period"), BUDGETS_SHEET, row_number, name)
                amount = _require_positive_float(row, "amount", BUDGETS_SHEET, row_number, name)
                budget_key = (owner.id, name, year, int(period))
                if budget_key in budgets_by_key:
                    raise ImportValidationError(
                        sheet_name=BUDGETS_SHEET,
                        row_number=row_number,
                        field_name="name",
                        message="Budget already exists",
                        raw_key=name,
                    )
                budget = Budget(
                    owner_id=owner.id,
                    name=name,
                    year=year,
                    period=int(period),
                    amount_cents=amount_to_cents(amount),
                )
                session.add(budget)
                session.commit()
                session.refresh(budget)
                budgets_by_key[budget_key] = budget
                job.success_rows += 1
            except ImportValidationError as exc:
                _record_job_error(session, job.id, exc)
                job.failed_rows += 1
            processed += 1
            _update_job_progress(
                session,
                job,
                progress_percent=_progress_percent(processed, total_rows),
            )

        pending_transactions: list[PendingTransactionRow] = []
        for row_number, row in _iter_sheet_rows(
            workbook[TRANSACTIONS_SHEET],
            TRANSACTION_HEADERS,
        ):
            try:
                owner = _require_user(
                    users_by_login, row["owner_login_name"], TRANSACTIONS_SHEET, row_number
                )
                category_name = _require_text(
                    row, "category_name", TRANSACTIONS_SHEET, row_number
                )
                category = categories_by_key.get((owner.id, category_name))
                if category is None:
                    raise ImportValidationError(
                        sheet_name=TRANSACTIONS_SHEET,
                        row_number=row_number,
                        field_name="category_name",
                        message="Category not found",
                        raw_key=category_name,
                    )
                budget_name = _optional_text(row.get("budget_name"))
                budget = None
                if budget_name:
                    budget = _resolve_budget_by_name(
                        budgets_by_key, owner.id, budget_name, TRANSACTIONS_SHEET, row_number
                    )
                handler_login_name = _optional_text(row.get("handler_login_name"))
                handler_user = (
                    _require_user(
                        users_by_login,
                        handler_login_name,
                        TRANSACTIONS_SHEET,
                        row_number,
                        field_name="handler_login_name",
                    )
                    if handler_login_name
                    else owner
                )
                summary = _require_text(row, "summary", TRANSACTIONS_SHEET, row_number)
                detail = _build_transaction_detail(
                    row.get("detail_note"), row.get("detail_items")
                )
                transaction = Transaction(
                    owner_id=owner.id,
                    category_id=category.id,
                    budget_id=budget.id if budget else None,
                    transaction_date=_parse_date(
                        row.get("transaction_date"),
                        TRANSACTIONS_SHEET,
                        row_number,
                        category_name,
                    ),
                    transaction_type=int(
                        _parse_transaction_type(
                            row.get("transaction_type"), TRANSACTIONS_SHEET, row_number, category_name
                        )
                    ),
                    amount_cents=amount_to_cents(
                        _require_positive_float(
                            row, "amount", TRANSACTIONS_SHEET, row_number, category_name
                        )
                    ),
                    entry_status=int(
                        _parse_entry_status(
                            row.get("entry_status"), TRANSACTIONS_SHEET, row_number, category_name
                        )
                    ),
                    handler_user_id=handler_user.id,
                    handler_name=handler_user.full_name or handler_user.login_name,
                    summary=summary,
                    description=summary,
                    detail=detail,
                )
                pending_transactions.append(
                    PendingTransactionRow(
                        transaction=transaction,
                        row_number=row_number,
                        raw_key=summary or category_name,
                    )
                )
                if len(pending_transactions) >= JOB_BATCH_SIZE:
                    _flush_transaction_batch(session, job, pending_transactions)
                    pending_transactions = []
            except ImportValidationError as exc:
                _record_job_error(session, job.id, exc)
                job.failed_rows += 1
            processed += 1
            _update_job_progress(
                session,
                job,
                progress_percent=_progress_percent(processed, total_rows),
            )
        if pending_transactions:
            _flush_transaction_batch(session, job, pending_transactions)

        if job.failed_rows > 0:
            _write_error_workbook(session, job)
        job.progress_percent = 100
        job.status = (
            DataJobStatus.PARTIAL_SUCCESS if job.failed_rows > 0 else DataJobStatus.SUCCEEDED
        )
        job.finished_at = _utcnow()
        session.add(job)
        session.commit()


def _flush_transaction_batch(
    session: Session, job: DataJob, batch: list[PendingTransactionRow]
) -> None:
    try:
        for pending_row in batch:
            session.add(pending_row.transaction)
        session.commit()
        job.success_rows += len(batch)
    except Exception as exc:  # noqa: BLE001
        session.rollback()
        for pending_row in batch:
            try:
                session.add(pending_row.transaction)
                session.commit()
                job.success_rows += 1
            except Exception as inner_exc:  # noqa: BLE001
                session.rollback()
                _record_job_error(
                    session,
                    job.id,
                    ImportValidationError(
                        sheet_name=TRANSACTIONS_SHEET,
                        row_number=pending_row.row_number,
                        field_name="row",
                        message=str(inner_exc or exc),
                        raw_key=pending_row.raw_key,
                    ),
                )
                job.failed_rows += 1
    finally:
        session.add(job)
        session.commit()


def _record_job_error(session: Session, job_id: uuid.UUID, exc: ImportValidationError) -> None:
    job_error = DataJobError(
        job_id=job_id,
        sheet_name=exc.sheet_name,
        row_number=exc.row_number,
        field_name=exc.field_name,
        error_message=exc.message,
        raw_key=exc.raw_key,
    )
    session.add(job_error)
    session.commit()


def _write_error_workbook(session: Session, job: DataJob) -> None:
    errors = session.exec(
        select(DataJobError)
        .where(DataJobError.job_id == job.id)
        .order_by(DataJobError.sheet_name, DataJobError.row_number)
    ).all()
    if not errors:
        return
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Errors"
    sheet.append(["sheet_name", "row_number", "field_name", "error_message", "raw_key"])
    for error in errors:
        sheet.append(
            [
                error.sheet_name,
                error.row_number,
                error.field_name,
                error.error_message,
                error.raw_key or "",
            ]
        )
    error_path = ensure_job_storage_dir() / f"errors-{job.id}.xlsx"
    workbook.save(error_path)
    job.error_file_path = str(error_path)
    session.add(job)
    session.commit()


def _populate_readme_sheet(sheet: Any) -> None:
    rows = [
        ["Template Version", TEMPLATE_VERSION],
        ["Sheet", "Description"],
        [CATEGORIES_SHEET, "owner_login_name, name, parent_name, color"],
        [BUDGETS_SHEET, "owner_login_name, name, year, period, amount"],
        [
            TRANSACTIONS_SHEET,
            ", ".join(TRANSACTION_HEADERS),
        ],
        ["Allowed period values", "MONTH | QUARTER | YEAR"],
        ["Allowed transaction_type values", "INCOME | EXPENSE"],
        ["Allowed entry_status values", "PENDING | SKIPPED | ENTERED"],
    ]
    for row in rows:
        sheet.append(row)


def _validate_required_sheets(sheet_names: list[str]) -> None:
    required = {README_SHEET, CATEGORIES_SHEET, BUDGETS_SHEET, TRANSACTIONS_SHEET}
    missing = required.difference(sheet_names)
    if missing:
        raise RuntimeError(f"Missing required sheets: {', '.join(sorted(missing))}")


def _read_template_version(sheet: Any) -> str | None:
    version_label = sheet["A1"].value
    if version_label != "Template Version":
        return None
    value = sheet["B1"].value
    return str(value).strip() if value is not None else None


def _iter_sheet_rows(sheet: Any, expected_headers: list[str]):
    rows = sheet.iter_rows(values_only=True)
    header_row = next(rows, None)
    if header_row is None:
        return
    headers = [str(value).strip() if value is not None else "" for value in header_row]
    if headers != expected_headers:
        raise RuntimeError(
            f"Invalid headers for {sheet.title}: expected {expected_headers}, got {headers}"
        )
    for index, values in enumerate(rows, start=2):
        record = {
            expected_headers[position]: values[position] if position < len(values) else None
            for position in range(len(expected_headers))
        }
        if all(value in (None, "") for value in record.values()):
            continue
        yield index, record


def _require_user(
    users_by_login: dict[str, User],
    login_name_value: Any,
    sheet_name: str,
    row_number: int,
    *,
    field_name: str = "owner_login_name",
) -> User:
    login_name = _optional_text(login_name_value)
    if not login_name or login_name not in users_by_login:
        raise ImportValidationError(
            sheet_name=sheet_name,
            row_number=row_number,
            field_name=field_name,
            message="User not found",
            raw_key=login_name,
        )
    return users_by_login[login_name]


def _require_text(row: dict[str, Any], field_name: str, sheet_name: str, row_number: int) -> str:
    value = _optional_text(row.get(field_name))
    if not value:
        raise ImportValidationError(
            sheet_name=sheet_name,
            row_number=row_number,
            field_name=field_name,
            message="Field is required",
        )
    return value


def _optional_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _detail_note(detail: dict[str, Any] | None) -> str:
    if not detail:
        return ""
    note = detail.get("note")
    return note if isinstance(note, str) else ""


def _serialize_detail_items(detail: dict[str, Any] | None) -> str:
    if not detail:
        return ""
    items = detail.get("items")
    if not isinstance(items, list):
        return ""
    lines: list[str] = []
    for item in items:
        if not isinstance(item, dict):
            continue
        name = _optional_text(item.get("name"))
        if not name:
            continue
        remark = _optional_text(item.get("remark"))
        lines.append(f"{name} - {remark}" if remark else name)
    return "\n".join(lines)


def _build_transaction_detail(note_value: Any, items_value: Any) -> dict[str, Any] | None:
    note = _optional_text(note_value)
    items_text = _optional_text(items_value)
    items: list[dict[str, Any]] = []
    if items_text:
        for raw_line in items_text.splitlines():
            line = raw_line.strip()
            if not line:
                continue
            if " - " in line:
                name, remark = line.split(" - ", 1)
                item: dict[str, Any] = {"name": name.strip()}
                remark = remark.strip()
                if remark:
                    item["remark"] = remark
                items.append(item)
            else:
                items.append({"name": line})
    if not note and not items:
        return None
    detail: dict[str, Any] = {}
    if note:
        detail["note"] = note
    if items:
        detail["items"] = items
    return detail


def _require_int(row: dict[str, Any], field_name: str, sheet_name: str, row_number: int) -> int:
    value = row.get(field_name)
    try:
        return int(value)
    except (TypeError, ValueError):
        raise ImportValidationError(
            sheet_name=sheet_name,
            row_number=row_number,
            field_name=field_name,
            message="Invalid integer value",
        )


def _require_positive_float(
    row: dict[str, Any],
    field_name: str,
    sheet_name: str,
    row_number: int,
    raw_key: str | None,
) -> float:
    value = row.get(field_name)
    try:
        amount = float(value)
    except (TypeError, ValueError):
        raise ImportValidationError(
            sheet_name=sheet_name,
            row_number=row_number,
            field_name=field_name,
            message="Invalid number",
            raw_key=raw_key,
        )
    if amount <= 0:
        raise ImportValidationError(
            sheet_name=sheet_name,
            row_number=row_number,
            field_name=field_name,
            message="Amount must be greater than 0",
            raw_key=raw_key,
        )
    return amount


def _parse_budget_period(
    value: Any, sheet_name: str, row_number: int, raw_key: str | None
) -> BudgetPeriod:
    text = _optional_text(value)
    try:
        if text and text.isdigit():
            return BudgetPeriod(int(text))
        if text:
            return BudgetPeriod[text.upper()]
    except (KeyError, ValueError):
        pass
    raise ImportValidationError(
        sheet_name=sheet_name,
        row_number=row_number,
        field_name="period",
        message="Invalid budget period",
        raw_key=raw_key,
    )


def _parse_transaction_type(
    value: Any, sheet_name: str, row_number: int, raw_key: str | None
) -> TransactionType:
    text = _optional_text(value)
    try:
        if text and text.isdigit():
            return TransactionType(int(text))
        if text:
            return TransactionType[text.upper()]
    except (KeyError, ValueError):
        pass
    raise ImportValidationError(
        sheet_name=sheet_name,
        row_number=row_number,
        field_name="transaction_type",
        message="Invalid transaction type",
        raw_key=raw_key,
    )


def _parse_entry_status(
    value: Any, sheet_name: str, row_number: int, raw_key: str | None
) -> EntryStatus:
    text = _optional_text(value)
    try:
        if text and text.isdigit():
            return EntryStatus(int(text))
        if text:
            return EntryStatus[text.upper()]
    except (KeyError, ValueError):
        pass
    raise ImportValidationError(
        sheet_name=sheet_name,
        row_number=row_number,
        field_name="entry_status",
        message="Invalid entry status",
        raw_key=raw_key,
    )


def _parse_date(value: Any, sheet_name: str, row_number: int, raw_key: str | None) -> date:
    if isinstance(value, date):
        return value
    text = _optional_text(value)
    if text is None:
        raise ImportValidationError(
            sheet_name=sheet_name,
            row_number=row_number,
            field_name="transaction_date",
            message="Transaction date is required",
            raw_key=raw_key,
        )
    try:
        return date.fromisoformat(text)
    except ValueError:
        raise ImportValidationError(
            sheet_name=sheet_name,
            row_number=row_number,
            field_name="transaction_date",
            message="Invalid date format, expected YYYY-MM-DD",
            raw_key=raw_key,
        )


def _resolve_budget_by_name(
    budgets_by_key: dict[tuple[uuid.UUID, str, int, int], Budget],
    owner_id: uuid.UUID,
    budget_name: str,
    sheet_name: str,
    row_number: int,
) -> Budget:
    matches = [
        budget
        for (candidate_owner_id, candidate_name, _, _), budget in budgets_by_key.items()
        if candidate_owner_id == owner_id and candidate_name == budget_name
    ]
    if len(matches) == 1:
        return matches[0]
    if not matches:
        raise ImportValidationError(
            sheet_name=sheet_name,
            row_number=row_number,
            field_name="budget_name",
            message="Budget not found",
            raw_key=budget_name,
        )
    raise ImportValidationError(
        sheet_name=sheet_name,
        row_number=row_number,
        field_name="budget_name",
        message="Budget name is ambiguous for this owner",
        raw_key=budget_name,
    )


def _update_job_progress(
    session: Session,
    job: DataJob,
    *,
    progress_percent: int | None = None,
    total_rows: int | None = None,
) -> None:
    if progress_percent is not None:
        job.progress_percent = progress_percent
    if total_rows is not None:
        job.total_rows = total_rows
    session.add(job)
    session.commit()


def _progress_percent(processed_rows: int, total_rows: int) -> int:
    if total_rows <= 0:
        return 100
    return min(99, int(processed_rows * 100 / total_rows))


def _mark_job_failed(job_id: uuid.UUID, reason: str) -> None:
    with Session(engine) as session:
        job = session.get(DataJob, job_id)
        if not job:
            return
        job.status = DataJobStatus.FAILED
        job.failure_reason = reason[:500]
        job.finished_at = _utcnow()
        session.add(job)
        session.commit()


def _utcnow():
    from app.models import get_datetime_utc

    return get_datetime_utc()
